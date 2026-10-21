"""Render the .pptx to PDF + PNG previews without Office/LibreOffice: every text box and picture is drawn at its
real position/size from the presentation (python-pptx -> matplotlib). Also flags text that overflows its box."""
from __future__ import annotations

import io
import textwrap

import matplotlib
from matplotlib.backends.backend_pdf import PdfPages
from pptx import Presentation
from pptx.util import Emu

from assetlens.deck.build_deck import OUT

matplotlib.use("Agg")
import matplotlib.image as mpimg  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402

PDF = OUT.with_suffix(".pdf")
EMU_IN = 914400
CHAR_W = 0.52   # average glyph width as a fraction of font size (sans-serif)


def _lines(text: str, size_pt: float, width_in: float, bold: bool = False) -> list[str]:
    cpl = max(10, int(width_in * 72 / (size_pt * (0.62 if bold else CHAR_W))))
    out = []
    for para in text.split("\n"):
        out += textwrap.wrap(para, cpl) or [""]
    return out


def render() -> dict:
    prs = Presentation(OUT)
    W, H = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    overflow = []
    with PdfPages(PDF) as pdf:
        for n, slide in enumerate(prs.slides, 1):
            fig = plt.figure(figsize=(W, H), dpi=110)
            ax = fig.add_axes([0, 0, 1, 1])
            ax.set_xlim(0, W); ax.set_ylim(H, 0); ax.axis("off")
            for sh in slide.shapes:
                x, y, w, h = (Emu(v).inches for v in (sh.left, sh.top, sh.width, sh.height))
                if sh.shape_type == 13:   # picture
                    img = mpimg.imread(io.BytesIO(sh.image.blob), format="png")
                    ax.imshow(img, extent=(x, x + w, y + h, y), aspect="auto", zorder=1)
                elif sh.has_text_frame:
                    cy = y
                    for p in sh.text_frame.paragraphs:
                        size = p.font.size.pt if p.font.size else 12
                        bold = bool(p.font.bold)
                        color = tuple(c / 255 for c in p.font.color.rgb) if p.font.color and p.font.color.type else (0, 0, 0)
                        for ln in _lines(p.text, size, w, bold):
                            ax.text(x, cy, ln.replace("$", "\\$"), fontsize=size, fontweight="bold" if bold else "normal", color=color, va="top", ha="left", zorder=2)
                            cy += size * 1.25 / 72
                        cy += 5 / 72
                    if cy > y + h + 0.05:
                        overflow.append((n, round(cy - (y + h), 2)))
            pdf.savefig(fig)
            fig.savefig(OUT.parent / f"slide_{n}.png")
            plt.close(fig)
    return dict(pdf=str(PDF), slides=len(prs.slides), overflow_in=overflow)


if __name__ == "__main__":
    print(render())
