import re

import pytest

from assetlens.deck.build_deck import OUT, build, numbers
from assetlens.warehouse.build import DB


@pytest.mark.skipif(not DB.exists(), reason="warehouse not built")
def test_every_headline_number_matches_warehouse():
    build()
    from pptx import Presentation
    text = " ".join(sh.text_frame.text for sl in Presentation(OUT).slides for sh in sl.shapes if sh.has_text_frame)
    n = numbers()   # re-derived independently from the warehouse
    must = [f"{n['GS_aum_tn']:.2f}", f"{n['GS_fee_yield']:.1f}", f"{n['BLK_fee_yield']:.1f}", f"{n['base_margin8']:.1f}%",
            f"{n['stress_margin8']:.1f}%", f"{n['be_marginal']:.1f}", f"{n['be_full']:.1f}", f"${n['gcc_npv']:.0f}m"]
    missing = [m for m in must if m not in text]
    assert not missing, missing
    assert len(re.findall("Nothing here is investment advice", text)) >= 1
    assert len(__import__("pptx").Presentation(OUT).slides) == 6


@pytest.mark.skipif(not DB.exists(), reason="warehouse not built")
def test_pdf_rendered_without_overflow():
    from assetlens.deck import render_pdf
    r = render_pdf.render()
    assert r["slides"] == 6 and not r["overflow_in"]
    assert render_pdf.PDF.exists() and render_pdf.PDF.stat().st_size > 50_000
